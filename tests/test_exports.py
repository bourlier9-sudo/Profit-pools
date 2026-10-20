"""Exports (Phase 7): every file is valid and parses back to the same numbers the workspace shows."""

import io
import json
import re
import zipfile

import numpy as np
import pandas as pd
import pytest

from profit_pools.core.config import CONFIG_DIR
from profit_pools.core.data.adapters import SyntheticAdapter
from profit_pools.core.model import views as V
from profit_pools.core.model.cube import build_history
from profit_pools.core.model.engine import run_baseline
from profit_pools.core.model.scenario import Scenario
from profit_pools.core.model.scenario_engine import run_scenario
from profit_pools.core.model.views import Geo
from profit_pools.export import FICTIONAL, ExportContext, csv_export, html_snapshot, pptx_summary
from profit_pools.export import figures as F


@pytest.fixture(scope="module")
def setup():
    a = SyntheticAdapter()
    m = a.metadata
    base = run_baseline(a)
    sc = Scenario.from_json((CONFIG_DIR / "scenarios" / "illustrative.json").read_text())
    res = run_scenario(build_history(a), m, sc).result
    return m, base, res, sc


def ctx_for(setup, view="compare", dataset="MARKET", year=2028, geo="US", focus=None):
    m, base, res, sc = setup
    return ExportContext(m, base if view == "baseline" else res, base, view, dataset, year, Geo.parse(m, geo),
                         focus=focus, scenario=None if view == "baseline" else sc)


@pytest.mark.parametrize("view", ["baseline", "compare", "impact"])
@pytest.mark.parametrize("focus", [None, "PROVIDER"])
def test_analysis_csv_matches_rollup(setup, view, focus):
    ctx = ctx_for(setup, view, focus=focus)
    df = pd.read_csv(io.BytesIO(csv_export.analysis_csv(ctx)))
    r = V.rollup(ctx.result, ctx.meta, ctx.dataset, ctx.year, focus, ctx.geo).rows
    assert list(df["pool_id"]) == list(r["node_id"])
    assert np.allclose(df["revenue_usd"], r["revenue"])
    assert ("delta_revenue_usd" in df) == (view != "baseline")
    if view != "baseline":
        b = V.rollup(ctx.base, ctx.meta, ctx.dataset, ctx.year, focus, ctx.geo).rows
        assert np.allclose(df["delta_revenue_usd"], r["revenue"].to_numpy() - b["revenue"].to_numpy())


def test_unavailable_stays_empty_not_zero(setup):
    df = pd.read_csv(io.BytesIO(csv_export.analysis_csv(ctx_for(setup, "baseline"))))
    assert df.loc[df["pool_id"] == "MANUFACTURER", "enrollment"].isna().all()


@pytest.mark.parametrize("view", ["baseline", "compare"])
def test_scenario_zip_round_trip(setup, view):
    ctx = ctx_for(setup, view)
    z = zipfile.ZipFile(io.BytesIO(csv_export.scenario_zip(ctx)))
    names = set(z.namelist())
    assert {"pools.csv", "flows.csv", "ownership.csv", "README.txt"} <= names
    assert ("scenario.json" in names) == (view != "baseline")
    pools = pd.read_csv(z.open("pools.csv"))
    rev = ctx.base.pools.groupby("dataset")["revenue"].sum()
    assert np.allclose(pools.groupby("dataset")["baseline_revenue"].sum(), rev)
    flows = pd.read_csv(z.open("flows.csv"))
    assert flows["baseline_allowed"].sum() == pytest.approx(ctx.base.flows["amount"].sum())
    assert FICTIONAL in z.read("README.txt").decode()
    if view != "baseline":
        assert np.allclose(pools["delta_revenue"].fillna(0), (pools["scenario_revenue"] - pools["baseline_revenue"]).fillna(0))
        s = ctx.result.pools.groupby("dataset")["revenue"].sum()
        assert np.allclose(pools.groupby("dataset")["scenario_revenue"].sum(), s)
        assert Scenario.from_json(z.read("scenario.json").decode()).to_json() == ctx.scenario.to_json()
        json.loads(z.read("scenario.json"))
        assert "credential" not in z.read("scenario.json").decode().lower()


@pytest.mark.parametrize("view, dataset, geo", [("compare", "MARKET", "US"), ("baseline", "ENTERPRISE", "NVL"),
                                                ("impact", "MARKET", "NVL-01")])
def test_html_snapshot_is_self_contained(setup, view, dataset, geo):
    html = html_snapshot.build(ctx_for(setup, view, dataset, geo=geo))
    assert html.startswith("<!doctype html>") and html.rstrip().endswith("</html>")
    assert not re.search(r"<script[^>]+src=", html)           # nothing loaded from the network
    assert html.count("plotly.js v") >= 1 and html.count("Plotly.newPlot") >= 3
    assert FICTIONAL in html and "not a simulator" in html
    if dataset == "ENTERPRISE":
        assert "Consolidated gain" in html


def _titles(prs):
    out = []
    for sl in prs.slides:
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text]
        out.append(texts[0] if texts else "")
    return out


@pytest.mark.skipif(not pptx_summary.available(), reason="python-pptx not installed")
@pytest.mark.parametrize("view, dataset, focus", [("compare", "MARKET", None), ("baseline", "ENTERPRISE", None),
                                                  ("impact", "MARKET", "PROVIDER")])
def test_pptx_opens_with_native_marimekko(setup, view, dataset, focus):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from profit_pools.export import headlines as HL
    ctx = ctx_for(setup, view, dataset, focus=focus)
    prs = Presentation(io.BytesIO(pptx_summary.build(ctx)))
    titles = _titles(prs)
    assert titles[1] == "Executive summary"
    mekko = list(prs.slides)[titles.index(HL.mekko(ctx))]
    rects = [s for s in mekko.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    tiles = pptx_summary._fold(F.tiles(ctx)[0])
    legend = 0 if view == "impact" else 6
    assert len(rects) == len(tiles) + legend                   # one native shape per (folded) Marimekko tile
    charts = [s for sl in prs.slides for s in sl.shapes if s.has_chart]
    assert len(charts) >= 2
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert FICTIONAL in text
    assert any(t.startswith("Owned assets add") for t in titles) == (dataset == "ENTERPRISE")
    assert any(ctx.scenario is not None and t.startswith(ctx.scenario.name) for t in titles) == (view != "baseline")
    summary = " ".join(sh.text_frame.text for sh in list(prs.slides)[1].shapes if sh.has_text_frame)
    for t in titles[2:-1]:
        assert t in summary                                   # the executive summary lists every headline


def test_fold_keeps_column_totals(setup):
    ctx = ctx_for(setup, "baseline")
    rows = F.tiles(ctx)[0]
    folded = pptx_summary._fold(rows)
    assert np.allclose(folded.groupby("col_id")["revenue"].sum(), rows.groupby("col_id")["revenue"].sum())
    assert np.allclose(folded.groupby("col_id")["share_in_col"].sum(), 1.0)


@pytest.mark.skipif(not pptx_summary.available(), reason="python-pptx not installed")
def test_pptx_uses_a_template(setup):
    from pptx import Presentation
    from pptx.util import Inches
    t = Presentation()
    t.slide_width, t.slide_height = Inches(10), Inches(7.5)           # 4:3 template with a sample slide
    t.slides.add_slide(t.slide_layouts[0]).shapes.title.text = "SAMPLE"
    buf = io.BytesIO()
    t.save(buf)
    assert pptx_summary.check_template(buf.getvalue()) is None
    assert pptx_summary.check_template(b"not a pptx")
    ctx = ctx_for(setup, "compare")
    plain = Presentation(io.BytesIO(pptx_summary.build(ctx)))
    tpl = Presentation(io.BytesIO(pptx_summary.build(ctx, buf.getvalue())))
    assert len(tpl.slides) == len(plain.slides)                   # sample slide removed
    assert tpl.slide_width == Inches(10)                           # template size kept
    assert "SAMPLE" not in " ".join(sh.text_frame.text for sl in tpl.slides for sh in sl.shapes if sh.has_text_frame)
    assert list(tpl.slides)[2].shapes.title.text == _titles(plain)[2]    # headline in the template's title box
