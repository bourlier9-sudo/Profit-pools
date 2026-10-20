"""PPTX summary with native PowerPoint objects (editable): KPI boxes, a Marimekko built from shapes,
native line and bar charts, and tables. No image rendering, so no Kaleido or Chrome dependency.
The module is optional: `available()` is False when python-pptx is not installed and the UI hides it."""

from __future__ import annotations

import io

import numpy as np

from profit_pools.charts import trends
from profit_pools.charts.marimekko import tile_fill
from profit_pools.core.model import flows_view as FV
from profit_pools.core.model import views as V
from profit_pools.core.model.scenario_engine import summarize
from profit_pools.design import format as fmt
from profit_pools.design.tokens import COLOR, LAYER_COLOR, OTHER, text_on
from profit_pools.export import FICTIONAL, ExportContext
from profit_pools.export import figures as F

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_PATTERN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
    _OK = True
except ImportError:  # pragma: no cover - optional dependency
    _OK = False

FONT = "Calibri"
W, H = 13.333, 7.5


def available() -> bool:
    return _OK


def _rgb(hex_color: str):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _text(slide, x, y, w, h, text, size=12, bold=False, color=None, align=None, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    if anchor:
        tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = _rgb(color or COLOR["ink_1"])
        if align:
            p.alignment = align
    return tb


def _slide(prs, ctx: ExportContext, title: str, subtitle: str | None = None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, 0.5, 0.35, W - 1, 0.5, title, 22, True)
    _text(s, 0.5, 0.85, W - 1, 0.35, subtitle or ctx.title(), 11, color=COLOR["ink_3"])
    line = s.shapes.add_connector(1, Inches(0.5), Inches(1.25), Inches(W - 0.5), Inches(1.25))
    line.line.color.rgb = _rgb(COLOR["line"])
    line.line.width = Pt(0.75)
    _text(s, 0.5, H - 0.4, W - 1, 0.3, FICTIONAL, 9, color=COLOR["ink_3"])
    return s


def _table(slide, x, y, w, head, rows, col_w=None, size=10, numeric_from=1):
    shape = slide.shapes.add_table(len(rows) + 1, len(head), Inches(x), Inches(y), Inches(w),
                                   Inches(0.3 * (len(rows) + 1)))
    t = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for j, h in enumerate(head):
        _cell(t.cell(0, j), h, size - 1, True, COLOR["ink_3"], j >= numeric_from)
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            _cell(t.cell(i, j), v, size, False, COLOR["ink_1"], j >= numeric_from)
    return t


def _cell(cell, text, size, bold, color, right):
    cell.fill.background()
    cell.margin_left = cell.margin_right = Inches(0.05)
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = str(text)
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = _rgb(color)


def _kpis(slide, ctx: ExportContext, y: float) -> None:
    items = F.kpis(ctx)
    if not items:
        return
    cw = (W - 1) / len(items)
    for i, (label, value, d) in enumerate(items):
        x = 0.5 + i * cw
        _text(slide, x, y, cw - 0.1, 0.25, label.upper(), 9, color=COLOR["ink_3"])
        _text(slide, x, y + 0.22, cw - 0.1, 0.45, value, 22, True)
        _text(slide, x, y + 0.7, cw - 0.1, 0.3, d, 9, color=COLOR["ink_2"])


def _marimekko(slide, ctx: ExportContext, x: float, y: float, w: float, h: float) -> None:
    rows, labels, est = F.tiles(ctx)
    if rows.empty:
        _text(slide, x, y, w, 0.4, "No pools to show for this view.", 12, color=COLOR["ink_3"])
        return
    color_by = "impact" if ctx.view == "impact" else "margin"
    scale = float(rows["d_gain"].abs().max()) if color_by == "impact" else 0.0
    cols = rows.drop_duplicates("col_id")[["col_id", "col_share"]]
    x0 = x
    gap = 0.03
    for c in cols.itertuples():
        cw = c.col_share * w
        if cw >= 0.5:
            _text(slide, x0, y - 0.32, cw, 0.3, labels[c.col_id] + (f"  {fmt.pct(c.col_share, 0)}" if cw > 1.2 else ""),
                  10, True, align=PP_ALIGN.CENTER)
        top = y
        for r in rows[rows["col_id"] == c.col_id].itertuples():
            th = r.share_in_col * h
            fill, affected = tile_fill(r, color_by, "d_gain", scale, V.IMPACT_MIN)
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(top),
                                         Inches(max(cw - gap, 0.01)), Inches(max(th - gap, 0.01)))
            shp.line.fill.background()
            if color_by != "impact" and r.revenue_with_margin <= 0:      # missing margin → hatch
                shp.fill.patterned()
                shp.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
                shp.fill.fore_color.rgb = _rgb(COLOR["missing_hatch"])
                shp.fill.back_color.rgb = _rgb(COLOR["missing"])
            else:
                shp.fill.solid()
                shp.fill.fore_color.rgb = _rgb(fill)
            if cw >= 0.7 and th >= 0.35:
                tf = shp.text_frame
                tf.margin_left = tf.margin_right = Inches(0.03)
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                val = (fmt.money(r.d_gain, signed=True) if affected else "no change") if color_by == "impact" \
                    else fmt.money(r.revenue)
                lab = r.seg_label + (" · est." if c.col_id in est else "")
                ink = _rgb(text_on(fill) if color_by != "impact" or affected else COLOR["ink_3"])
                for k, line in enumerate([lab] if th < 0.6 else [lab, val]):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = line
                    run.font.size, run.font.name = Pt(9), FONT
                    run.font.color.rgb = ink
            top += th
        x0 += cw
    legend = ("Color = change in operating gain vs baseline (teal up, red down; grey = not affected)."
              if color_by == "impact" else
              "Width = share of revenue · height = share within column · color = operating margin "
              "(red < 0%, grey 0%, teal > 0%) · hatched = margin unavailable · est. = top-down estimate.")
    _text(slide, x, y + h + 0.05, w, 0.3, legend, 9, color=COLOR["ink_3"])


def _line_chart(slide, ctx: ExportContext, x, y, w, h) -> None:
    df = F.trend_frame(ctx)
    if df.empty:
        return
    years = sorted(df["year"].unique())
    cd = CategoryChartData()
    cd.categories = [str(int(yr)) for yr in years]
    ids = list(dict.fromkeys(df["node_id"]))
    for nid in ids:
        s = df[df["node_id"] == nid].set_index("year")["revenue"].reindex(years) / 1e9
        cd.add_series(ctx.meta.taxonomy[nid].short_label, [None if not np.isfinite(v) else round(float(v), 2) for v in s])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = _rgb(COLOR["line"])
    ch.value_axis.tick_labels.number_format = '"$"#,##0"B"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.font.size = Pt(9)
    ch.category_axis.tick_labels.font.size = Pt(9)
    colors = trends.series_colors(ids, LAYER_COLOR if ctx.focus is None else {})
    for i, nid in enumerate(ids):
        ser = ch.series[i]
        ser.format.line.width = Pt(2)
        ser.format.line.color.rgb = _rgb(colors.get(nid, OTHER))
        ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE


def _bar_chart(slide, labels, values, x, y, w, h, number_format='"$"#,##0.0"B"') -> None:
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("Change", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.number_format = number_format
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(9)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.tick_labels.font.size = Pt(9)
    ch.value_axis.tick_labels.number_format = number_format
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.has_major_gridlines = False
    ser = ch.series[0]
    for i, v in enumerate(values):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(COLOR["accent"] if v >= 0 else COLOR["bad"])


def build(ctx: ExportContext) -> bytes:
    if not _OK:
        raise RuntimeError("python-pptx is not installed")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(Inches(W))), Emu(int(Inches(H)))
    m = ctx.meta

    # 1 — title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, 0.8, 2.4, W - 1.6, 0.8, "Profit Pools Explorer", 36, True)
    _text(s, 0.8, 3.3, W - 1.6, 0.5, ctx.title(), 16, color=COLOR["ink_2"])
    _text(s, 0.8, 4.1, W - 1.6, 0.4, FICTIONAL, 12, color=COLOR["ink_3"])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.9), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(COLOR["accent"])
    bar.line.fill.background()

    # 2 — KPIs + Marimekko
    s = _slide(prs, ctx, "Where revenue and operating gain sit")
    _kpis(s, ctx, 1.4)
    _marimekko(s, ctx, 0.5, 2.85, W - 1, 4.0)

    # 3 — revenue trend
    s = _slide(prs, ctx, "Revenue trend by pool", f"{ctx.scope_label()} · {ctx.geo_label()} · $B · forecast from "
                                                    f"{m.periods.base_year + 1}")
    _line_chart(s, ctx, 0.5, 1.45, W - 1, 5.4)

    # 4 — growth drivers
    br = F.growth_bridge(ctx)
    d = br[br["kind"] == "delta"]
    p = m.periods
    s = _slide(prs, ctx, f"What drives revenue growth, {p.base_year}–{p.forecast_end}",
               f"{ctx.scope_label()} · revenue {fmt.money(br['value'].iloc[0])} → {fmt.money(br['value'].iloc[-1])}")
    if len(d):
        _bar_chart(s, list(d["label"]), [round(float(v) / 1e9, 2) for v in d["value"]], 0.5, 1.45, W - 1, 5.2)
    _text(s, 0.5, 6.7, W - 1, 0.3, "Sequential split; components sum exactly to the change in revenue.", 9,
          color=COLOR["ink_3"])

    # 5 — value flows (largest links)
    view = FV.build(ctx.result, m, ctx.dataset, ctx.year, ctx.geo, FV.SankeyConfig())
    if len(view.links):
        nl = view.nodes.set_index("node_id")["label"]
        top = view.links.sort_values("value", ascending=False).head(14)
        s = _slide(prs, ctx, "Largest value flows", f"Allowed spending, {ctx.year} · {ctx.geo_label()} · total "
                                                     f"{fmt.money(view.total)}")
        _table(s, 0.5, 1.45, W - 1, ["From", "To", "Allowed spend", "% of total", "% of source"],
               [[nl[r.source], nl[r.target], fmt.money(r.value), fmt.pct(r.share_of_total), fmt.pct(r.share_of_source)]
                for r in top.itertuples()], col_w=[3.4, 3.7, 1.9, 1.6, 1.73], numeric_from=2)

    # 6 — scenario
    if ctx.has_scenario:
        s = _slide(prs, ctx, f"Scenario: {ctx.scenario.name}", ctx.scenario.description or ctx.title())
        rows = []
        for i, a in enumerate(ctx.scenario.enabled, start=1):
            rows.append([f"{i}. {a.label or ''}".strip(), summarize(m, a)])
        _table(s, 0.5, 1.45, W - 1, ["Assumption", "What changes"], rows, col_w=[3.0, W - 4.0], size=10,
               numeric_from=2)

    # 7 — value capture
    if ctx.dataset == "ENTERPRISE":
        v = V.ownership_view(ctx.result, m, ctx.year, ctx.geo)
        if not v.note:
            t = v.totals
            s = _slide(prs, ctx, "Value capture: owned assets and the consolidated P&L")
            _table(s, 0.5, 1.45, 5.2, ["", "Amount"], [
                ["Plan premium revenue", fmt.money(t["payer_revenue"])],
                ["+ Owned-asset revenue", fmt.money(t["owned_revenue"])],
                ["− Claims paid to own assets (eliminated)", fmt.money(-t["elimination"])],
                ["= Consolidated revenue", fmt.money(t["consolidated_revenue"])],
                ["Plan operating gain", fmt.money(t["payer_gain"])],
                ["+ Owned-asset gain", fmt.money(t["owned_gain"])],
                ["= Consolidated gain", fmt.money(t["consolidated_gain"])]], col_w=[3.6, 1.6])
            _table(s, 6.0, 1.45, W - 6.5, ["Asset", "Routed share", "Owned revenue", "Margin", "Owned gain"],
                   [[r.label, fmt.pct(r.routing_share), fmt.money(r.owned_revenue), fmt.pct(r.owned_margin),
                     fmt.money(r.owned_gain)] for r in v.assets.itertuples()], col_w=[2.6, 1.1, 1.2, 0.9, 1.03])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
